#!/usr/bin/env python3
"""Build the three department handouts from the adjacent Markdown sources.

Run with the Codex primary runtime Python (python-docx). Rendering/QA is a
separate step with the documents skill render_docx.py; no AutoCAD code runs.
"""
from pathlib import Path
import hashlib
import json
import re
import argparse

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
OUTPUTS = {
    "INSTALL_AUTOCAD_2024.md": "Install_AutoCAD_2024_build105.docx",
    "PILOT_GUIDE.md": "Facades_Pilot_build105.docx",
    "FEEDBACK_TEMPLATE.md": "Facades_Feedback_build105.docx",
}
LINKS = {**OUTPUTS, "Facades_Pilot_build105.docx": "Facades_Pilot_build105.docx"}


def element(tag, **attrs):
    e = OxmlElement(tag)
    for key, value in attrs.items():
        e.set(qn(key), str(value))
    return e


def clean_heading(text):
    text = re.sub(r"[`*_]", "", text)
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", text)).strip()


def new_document(title):
    d = Document()
    # The bundled Word template may carry a blue border in Title. Strip all
    # inherited paragraph borders so the handout uses typography only.
    for border in list(d.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.top_margin = s.bottom_margin = Inches(0.7)
    s.left_margin = s.right_margin = Inches(0.75)
    s.header_distance = s.footer_distance = Inches(0.3)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"):
        st = d.styles[name]
        st.font.name = "Arial"
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.size = Pt(11)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.07
        st.paragraph_format.widow_control = True
        st.element.get_or_add_rPr().append(element("w:lang", **{"w:val": "ru-RU"}))
    d.styles["Title"].font.size = Pt(20)
    d.styles["Title"].font.bold = True
    d.styles["Title"].paragraph_format.space_after = Pt(12)
    for n, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        st = d.styles[n]
        st.font.size = Pt(size)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
    footer = s.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("build 105  ·  ")
    r.font.name = "Arial"
    r.font.size = Pt(9)
    field = element("w:fldSimple", **{"w:instr": "PAGE"})
    rr = element("w:r")
    rr.append(element("w:t"))
    rr[-1].text = "1"
    field.append(rr)
    footer._p.append(field)
    d.core_properties.title = title
    d.core_properties.subject = "Ограниченный фасадный пилот build 105 для проектного отдела"
    d.core_properties.author = "AutoCAD Facades project"
    d.core_properties.keywords = "build-105, AutoCAD 2024, проверочный пилот"
    d.core_properties.comments = "Created from the adjacent versioned Markdown sources."
    return d


def plain_run(p, text, bold=False, italic=False, code=False):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    # Commands and paths stay at the same readable size as body text.
    if code:
        r.font.name = "Arial"
    return r


def inline(p, text):
    pattern = r"(\[[^\]]+\]\([^\)]+\)|\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)"
    for part in re.split(pattern, text):
        if not part:
            continue
        m = re.fullmatch(r"\[([^\]]+)\]\(([^\)]+)\)", part)
        if m:
            label, url = m.groups()
            label = label.replace("**", "").replace("`", "")
            url = LINKS.get(url, url)
            link = element("w:hyperlink", **{"r:id": p.part.relate_to(url, RT.HYPERLINK, is_external=True)})
            run = element("w:r")
            prop = element("w:rPr")
            prop.append(element("w:color", **{"w:val": "000000"}))
            prop.append(element("w:u", **{"w:val": "single"}))
            run.append(prop)
            t = element("w:t")
            t.text = label
            run.append(t)
            link.append(run)
            p._p.append(link)
        elif part.startswith("**"):
            plain_run(p, part[2:-2], bold=True)
        elif part.startswith("`"):
            plain_run(p, part[1:-1], code=True)
        elif part.startswith("*"):
            plain_run(p, part[1:-1], italic=True)
        else:
            plain_run(p, part)


def table(d, lines, feedback=False):
    rows = [[x.strip() for x in line.strip().strip("|").split("|")] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r"[:\- ]+", v or " ") for v in r)]
    if feedback:
        replacements = {
            "N1: Handle/блок/результат проверки": "N1\nHandle и блок\nПроверка",
            "N2: Handle/блок/результат проверки": "N2\nHandle и блок\nПроверка",
            "N1:Handle/блок/результат проверки": "N1\nHandle и блок\nПроверка",
            "N2:Handle/блок/результат проверки": "N2\nHandle и блок\nПроверка",
            "После сохранения/открытия": "После сохранения и открытия",
            "ATFTABLE: проверка/просмотр": "ATFTABLE:\nпроверка и просмотр",
        }
        rows = [[replacements.get(v, v) for v in r] for r in rows]
    cols = len(rows[0])
    # Field/value forms use ordinary editable paragraphs, not a dense grid.
    if feedback and rows[0] == ["Поле", "Заполнить"]:
        for label, _ in rows[1:]:
            p = d.add_paragraph()
            p.paragraph_format.keep_with_next = True
            inline(p, "**" + label + "**")
            p = d.add_paragraph("[Заполнить]")
            p.paragraph_format.space_after = Pt(8)
        return
    if cols == 2:
        widths = [2.1, 4.9]
    elif cols == 3:
        widths = [1.55, 2.7, 2.75]
    elif cols == 4:
        widths = [2.8, 1.4, 1.25, 1.55]
        if rows[0][0] == "Пункт":
            widths = [0.57, 3.10, 0.92, 2.41]
    elif cols == 5:
        widths = [1.4, 1.35, 1.6, 1.6, 1.05]
    elif cols == 6:
        widths = [1.15, 1.15, 1.2, 1, 1.25, 1.25]
    else:
        widths = [7 / cols] * cols
    t = d.add_table(rows=0, cols=cols)
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    pr = t._tbl.tblPr
    borders = element("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(element("w:" + side, **{"w:val": "single", "w:sz": "4", "w:color": "D9D9D9"}))
    pr.append(borders)
    margins = element("w:tblCellMar")
    for side, twips in (("top", 80), ("bottom", 80), ("left", 90), ("right", 90)):
        margins.append(element("w:" + side, **{"w:w": twips, "w:type": "dxa"}))
    pr.append(margins)
    for index, data in enumerate(rows):
        row = t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(element("w:cantSplit"))
        if index == 0:
            trpr.append(element("w:tblHeader"))
        for i, (cell, content) in enumerate(zip(row.cells, data)):
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            if not content and feedback:
                content = "[Заполнить]"
            inline(p, content)
            if index == 0:
                cell._tc.get_or_add_tcPr().append(element("w:shd", **{"w:fill": "EAEAEA"}))
                for r in p.runs:
                    r.bold = True
            elif rows[0][0] == "Пункт" and i in (0, 2):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 0.25
    p.add_run().font.size = Pt(2)


def markdown(d, text, feedback=False, title_override=None):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i]); i += 1
            table(d, block, feedback)
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            level, text = len(heading[1]), heading[2]
            if level == 1:
                d.add_paragraph(clean_heading(title_override or text), "Title")
            else:
                d.add_paragraph(clean_heading(text), "Heading " + str(min(level - 1, 3)))
            i += 1
            continue
        bullet = re.match(r"^(-|\d+\.)\s+(.+)$", line)
        block = [bullet[2] if bullet else line]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or re.match(r"^(#|\||- |\d+\. )", nxt):
                break
            block.append(nxt); i += 1
        text = " ".join(block)
        p = d.add_paragraph()
        if bullet:
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            plain_run(p, ("•" if bullet[1] == "-" else bullet[1]) + " ")
        inline(p, text)
        if feedback and "___" in text:
            p.paragraph_format.space_after = Pt(12)


def section(text, name):
    match = re.search(r"^## " + re.escape(name) + r"\s*$", text, re.M)
    if not match:
        raise ValueError(name)
    start = match.start()
    nxt = re.search(r"^## ", text[match.end():], re.M)
    return start, match.end() + nxt.start() if nxt else len(text)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--only", choices=("install", "pilot", "feedback"))
    args = parser.parse_args()
    chosen = {"install": "INSTALL_AUTOCAD_2024.md", "pilot": "PILOT_GUIDE.md", "feedback": "FEEDBACK_TEMPLATE.md"}.get(args.only)
    sources = {name: (ROOT / name).read_text(encoding="utf-8") for name in (*OUTPUTS, "ENGINEERING_QUESTIONS.md")}
    manifest = {"build": 105, "source_sha": "c4009838875aa0f225971153bf4b54fe288e8e79", "sources": {}, "outputs": {}, "editorial_changes": []}
    for name in sources:
        manifest["sources"][name] = hashlib.sha256((ROOT / name).read_bytes()).hexdigest()
    for source, output in OUTPUTS.items():
        if chosen and source != chosen:
            manifest["outputs"][output] = hashlib.sha256((ROOT / output).read_bytes()).hexdigest()
            continue
        d = new_document(output[:-5])
        content = sources[source]
        if source == "FEEDBACK_TEMPLATE.md":
            for lead in ("Лишние пустые схемы", "Статус определён новой командой", "Если требуются внутренние данные"):
                content = content.replace("\n" + lead, "\n\n" + lead)
        markdown(d, content, feedback=source == "FEEDBACK_TEMPLATE.md")
        if source == "INSTALL_AUTOCAD_2024.md":
            for p in d.paragraphs:
                if p.text == "Проверка установки":
                    p.paragraph_format.page_break_before = True
        if source == "FEEDBACK_TEMPLATE.md":
            for p in d.paragraphs:
                if p.text.startswith(("5 Карточка замечания", "6 Скорость на реальном фасаде", "7 Итог отдела")):
                    p.paragraph_format.page_break_before = True
        if source == "PILOT_GUIDE.md":
            questions = sources["ENGINEERING_QUESTIONS.md"]
            # The full pilot guide already explains each product boundary. The
            # appendix keeps all six questions and their engineering premises.
            start, end = section(questions, "Что именно проверяется в выпускаемом пилоте")
            questions = questions[:start] + "Программный маршрут и ограничения функций приведены в первой части этого документа. До фактического прохода отдела работа внутри AutoCAD остаётся непроверенной.\n\n" + questions[end:]
            start, end = section(questions, "Отдельно: сообщение об ошибке программы")
            questions = questions[:start] + "## Сообщение об ошибке программы\n\nИнженерные ответы и ошибки программы фиксируются отдельно. Для ошибки заполните карточку в Facades_Feedback_build105.docx и приложите материалы по П17. Предусмотренный отказ расчёта не означает дефект; потеря результата, неверные количества, положение таблицы, зависание и ошибка сохранения требуют воспроизведения.\n\n" + questions[end:]
            d.add_page_break()
            markdown(d, questions, title_override="Инженерные вопросы для завершения одного решения")
        d.save(ROOT / output)
        manifest["outputs"][output] = hashlib.sha256((ROOT / output).read_bytes()).hexdigest()
    manifest["editorial_changes"] = ["Combined PILOT_GUIDE.md with ENGINEERING_QUESTIONS.md.", "Replaced the duplicated product-status table with a reference to the full pilot guide.", "Replaced the duplicated defect-report table with a reference to the editable feedback form and П17; retained the distinction between expected engineering limits and defects.", "Sanitized title/heading punctuation for Word styles; preserved body meaning and all six engineering questions.", "Converted the passport field/value table into editable labelled paragraphs; all feedback route statuses remain NOT_RUN.", "Expanded compressed slash separators in narrow feedback cells and started the reusable defect card on a new page.", "Mapped installation, pilot and feedback Markdown hyperlinks to their deliverable DOCX names."]
    (ROOT / "document_sources.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
