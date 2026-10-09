"""Build the illustrated beginner handbook from the maintained Markdown source."""
import argparse
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

HERE = Path(__file__).resolve().parent
BLACK = RGBColor(0, 0, 0)


def plain(text):
    return re.sub(r'[`*]', '', text)


def title(text):
    return re.sub(r'\s+', ' ', re.sub(r'[^\w\s]', ' ', plain(text))).strip()


def hyperlink(p, label, target, internal=False):
    h = OxmlElement('w:hyperlink')
    if internal:
        h.set(qn('w:anchor'), target)
    else:
        h.set(qn('r:id'), p.part.relate_to(target, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r')
    pr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '245D83'); pr.append(color)
    r.append(pr)
    t = OxmlElement('w:t'); t.text = label; r.append(t); h.append(r); p._p.append(h)


def inline(p, text):
    parts = re.split(r'(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^\)]+\))', text)
    for token in parts:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            r = p.add_run(plain(token[2:-2])); r.bold = True
        elif token.startswith('`') and token.endswith('`'):
            r = p.add_run(token[1:-1]); r.font.name = 'DejaVu Sans Mono'; r.font.size = Pt(9.5)
        elif re.fullmatch(r'\[[^\]]+\]\([^\)]+\)', token):
            label, target = re.match(r'\[([^\]]+)\]\(([^\)]+)\)', token).groups()
            hyperlink(p, label, target)
        else:
            p.add_run(token)


def table(doc, rows):
    values = [[cell.strip() for cell in r.strip().strip('|').split('|')] for r in rows]
    values = [r for r in values if not all(re.fullmatch(r':?-+:?', c.replace(' ', '')) for c in r)]
    cols = max(len(r) for r in values)
    t = doc.add_table(rows=0, cols=cols)
    t.autofit = False
    lens = [max((len(plain(r[j])) for r in values if len(r) > j), default=1) for j in range(cols)]
    weights = [max(12, min(75, n)) ** .7 for n in lens]
    widths = [7.0 * w / sum(weights) for w in weights]
    if values[0] == ["Кнопка", "Команда", "Для чего нужна"]:
        widths = [1.6, 1.65, 3.75]  # Keep command names on one line.
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b = OxmlElement('w:' + edge); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4')
        b.set(qn('w:color'), 'D9D9D9'); borders.append(b)
    pr.append(borders)
    for i, row in enumerate(values):
        cells = t.add_row().cells
        trpr = cells[0]._tc.getparent().get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); trpr.append(cant)
        if i == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for j, cell in enumerate(cells):
            cell.width = Inches(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge, value in (('top', '85'), ('bottom', '85'), ('left', '100'), ('right', '100')):
                x = OxmlElement('w:' + edge); x.set(qn('w:w'), value); x.set(qn('w:type'), 'dxa'); margins.append(x)
            tcpr.append(margins)
            sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), '23485F' if i == 0 else ('F1F5F7' if i % 2 == 0 else 'FFFFFF')); tcpr.append(sh)
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            inline(p, row[j] if len(row) > j else '')
            for r in p.runs:
                r.font.size = Pt(10.2); r.font.color.rgb = RGBColor(255, 255, 255) if i == 0 else BLACK
                r.bold = i == 0 or r.bold
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', type=Path, default=HERE / 'Facades_User_Manual.docx')
    args = ap.parse_args()
    lines = (HERE / 'manual_content.md').read_text(encoding='utf-8').splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
    sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
    sec.header_distance = Inches(.25); sec.footer_distance = Inches(.3)
    for name in ('Normal', 'Body Text', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption'):
        style = doc.styles[name]
        style.font.name = 'DejaVu Sans'; style.font.color.rgb = BLACK
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.12
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    doc.styles['Title'].font.italic = False
    doc.styles['Subtitle'].font.italic = False
    doc.styles['Title'].font.size = Pt(26)
    doc.styles['Subtitle'].font.size = Pt(16)
    for level, size in ((1, 17), (2, 13), (3, 11.5)):
        st = doc.styles['Heading ' + str(level)]
        st.font.size = Pt(size); st.font.bold = True
        st.paragraph_format.space_before = Pt(16 if level == 1 else 10)
        st.paragraph_format.space_after = Pt(7); st.paragraph_format.keep_with_next = True
    doc.styles['Caption'].font.size = Pt(9)
    doc.styles['Caption'].paragraph_format.space_after = Pt(10)
    doc.core_properties.title = 'Фасадные модули AutoCAD Руководство начинающего пользователя'
    doc.core_properties.subject = 'Зоны облицовка подсистема и ведомости'
    doc.core_properties.author = 'Фасадные модули AutoCAD'
    doc.core_properties.keywords = 'AFacades AClad AFrame ATFZONE ATTILE ATFRAME ATFTABLE ATFPROJECT ATFZONEPARAMS ATFNODE ATFNODEIMPORT ATFNODEEDIT ATFNODEDEMO ATFNODETEST'
    head = sec.header.paragraphs[0]; head.text = 'Фасадные модули AutoCAD  •  Руководство начинающего пользователя'
    for r in head.runs: r.font.size = Pt(8); r.font.color.rgb = BLACK
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run(); r.font.size = Pt(9)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); r._r.addnext(field)

    headings = [(i, line[3:]) for i, line in enumerate(lines) if line.startswith('## ')]
    bookmark = {i: 'chapter_' + str(n + 1) for n, (i, _) in enumerate(headings)}
    hcount = 0; n = 0; in_toc = False; figure = 0
    while n < len(lines):
        line = lines[n].strip()
        if not line:
            n += 1; continue
        if line.startswith('# ') and not line.startswith('## '):
            hcount += 1
            p = doc.add_paragraph(title(line[2:]), style='Title' if hcount == 1 else 'Subtitle')
            n += 1; continue
        if line.startswith('## ') and not in_toc:
            p = doc.add_paragraph()
            r = p.add_run('Редакция от 9 октября 2026 года для выпуска build-111. Перед проверкой заново установите AFacades, AClad и AFrame из этого выпуска. Результаты проверки в AutoCAD фиксируйте отдельно для каждого сценария.')
            r.bold = True
            doc.add_page_break()
            doc.add_paragraph('Содержание', style='Heading 1')
            for i, text in headings:
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5)
                hyperlink(p, plain(text), bookmark[i], internal=True)
            doc.add_page_break(); in_toc = True
        if line.startswith('## '):
            p = doc.add_paragraph(title(line[3:]), style='Heading 1')
            if line.startswith('## 20.'):
                p.paragraph_format.page_break_before = True
            start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), str(n)); start.set(qn('w:name'), bookmark[n])
            end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(n)); p._p.insert(0, start); p._p.append(end)
            n += 1; continue
        if line.startswith('### '):
            doc.add_paragraph(title(line[4:]), style='Heading 2'); n += 1; continue
        if line.startswith('#### '):
            doc.add_paragraph(title(line[5:]), style='Heading 3'); n += 1; continue
        if line.startswith('|'):
            rows = []
            while n < len(lines) and lines[n].strip().startswith('|'):
                rows.append(lines[n]); n += 1
            table(doc, rows); continue
        im = re.fullmatch(r'!\[([^\]]*)\]\(([^\)]+)\)', line)
        if im:
            label, path = im.groups(); path = HERE / path
            w, h = Image.open(path).size
            width = min(7.0, 6.35 * w / h)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(path), width=Inches(width))
            for node in p._p.iter(qn('wp:docPr')):
                node.set('descr', label)
            figure += 1
            p = doc.add_paragraph(f'Рисунок {figure}. {label}', style='Caption')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            n += 1; continue
        if line.startswith('```'):
            n += 1
            while n < len(lines) and not lines[n].strip().startswith('```'):
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.left_indent = Inches(.12)
                r = p.add_run(lines[n]); r.font.name = 'DejaVu Sans Mono'; r.font.size = Pt(9)
                n += 1
            n += 1; continue
        match = re.match(r'^(\d+\.\s+|[-*]\s+)(.*)', line)
        if match:
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(.22)
            p.paragraph_format.first_line_indent = Inches(-.22)
            prefix = match.group(1) if match.group(1)[0].isdigit() else '• '
            inline(p, prefix + match.group(2)); n += 1; continue
        para = [line]; n += 1
        while n < len(lines) and lines[n].strip() and not re.match(r'^(#|\||!\[|```|\d+\.\s|[-*]\s)', lines[n].strip()):
            para.append(lines[n].strip()); n += 1
        inline(doc.add_paragraph(), ' '.join(para))
    args.out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.out)
    print(args.out)


if __name__ == '__main__':
    main()
