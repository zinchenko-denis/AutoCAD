"""Create the short native-node pilot handout. Private diagram is an input, never embedded in Git."""
import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor


RELEASE_ASSETS = (
    ("AFacades.bundle.zip", "Зоны и ведомости"),
    ("AClad.bundle.zip", "Облицовка"),
    ("AFrame.bundle.zip", "Подсистема и команды узлов"),
    ("Install_AutoCAD_2024.docx", "Скачивание, установка и обновление"),
    ("Facades_User_Manual.docx", "Полное руководство и контрольные фасады"),
)


def release_links(release_build):
    if release_build is None:
        return []
    if isinstance(release_build, bool) or not isinstance(release_build, int) or release_build < 110:
        raise ValueError("The node commands require an explicitly selected build 110 or newer")
    base = f"https://github.com/zinchenko-denis/AutoCAD/releases/download/build-{release_build}/"
    return [(name, purpose, base + name) for name, purpose in RELEASE_ASSETS]


def hyperlink(paragraph, text, url):
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    label = OxmlElement("w:t")
    label.text = text
    run.append(label)
    link.append(run)
    paragraph._p.append(link)


def paragraph(doc, text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    return p


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.style = "Table Grid"
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + edge)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn("w:" + key), value)
        borders.append(border)
    t._tbl.tblPr.append(borders)
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = text
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
    for row_index, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.space_before = Pt(4)
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    r.bold = row_index == 0
            if row_index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "DDEAF0")
                cell._tc.get_or_add_tcPr().append(shade)
    return t


def step(doc, number, text):
    p = paragraph(doc, f"{number}. {text}")
    p.paragraph_format.space_after = Pt(8)
    return p


def page(doc, title):
    doc.add_page_break()
    doc.add_heading(title, 1)


def build(out, diagram, code_sha, ci_url, release_build=None,
          fixes_name="Fixes_2026-10-08_09_build111.txt"):
    downloads = release_links(release_build)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.65), Cm(1.5)
    sec.left_margin, sec.right_margin = Cm(1.8), Cm(1.8)
    sec.header_distance = sec.footer_distance = Cm(.65)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"):
        style = doc.styles[name]
        style.font.name = "DejaVu Sans"
        style.font.color.rgb = RGBColor.from_string("000000")
        style.font.size = Pt(10.5)
        style.paragraph_format.line_spacing = 1.08
        style.paragraph_format.space_after = Pt(7)
    doc.styles["Title"].font.size = Pt(25)
    # Some runtime templates add a decorative border below the title.
    for border in doc.styles["Title"].element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    doc.styles["Subtitle"].font.size = Pt(12)
    doc.styles["Heading 1"].font.size = Pt(17)
    doc.styles["Heading 2"].font.size = Pt(12)
    doc.styles["Caption"].font.size = Pt(8.5)
    for name in ("Heading 1", "Heading 2"):
        doc.styles[name].paragraph_format.keep_with_next = True
    header = sec.header.paragraphs[0]
    header.text = "ФАСАДЫ / ПРОЕКТНАЯ ГРУППА / 09.10.2026"
    header.runs[0].font.size = Pt(8)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Тестовый узел · ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "Тестовый динамический узел — инструкция проектной группе"
    doc.core_properties.author = "Фасадный комплекс AutoCAD"
    doc.core_properties.subject = "Проверка выбранных трёх экземпляров из пяти в AutoCAD 2024"

    doc.add_heading("Тестовый\nдинамический узел", 0)
    doc.add_paragraph("Вектор-1 · тип 1 · узел 4.2.1\nИнструкция проектной группе · 9 октября 2026", "Subtitle")
    paragraph(doc, "Цель: изменить геометрию, размеры и марки трёх выбранных вставок из пяти. Две невыбранные вставки должны остаться прежними.", True)
    paragraph(doc, "В комплекте LISP-автор редакции 2026-10-08.2 для полного AutoCAD 2024 Windows. Исправлены пропущенный ответ о числе ручек Visibility, вторичная ошибка обработчика и учёт собственной группы Undo при исключении Begin/End. Готовый DWG ещё не создан; живая приёмка не завершена. Установка трёх фасадных модулей и проверка автора узла выполняются по отдельным шагам ниже.")
    table(doc, ["Параметр", "Состояние A", "Состояние B"], [
        ("Утеплитель", "100 мм", "150 мм"),
        ("Вынос до лицевой плоскости", "230 мм", "280 мм"),
        ("Кронштейн / профиль", "КР2-70-200 / ГП-40-40-1,2", "КР2-70-250 / ГП-60-40-1,2"),
        ("Удлинитель", "УК-70-100-1,2", "УК-70-100-1,2"),
    ], [6.0, 5.7, 5.7])
    if diagram:
        doc.add_picture(str(diagram), width=Cm(17.3))
        doc.add_paragraph("Предварительная схема из исходных координат. Это не снимок проверенного DWG.", "Caption")
    paragraph(doc, "Учебная геометрия: неразмеренные пятки и гибы условны; отверстия, кляммер и крепёж не моделируются. Вынос не подбирает кронштейн автоматически. Монтажная пригодность и несущая способность здесь не подтверждаются.")

    page(doc, "1 Создать тестовый DWG")
    paragraph(doc, "Эта часть выполняется без фасадных DLL. Она проверяет динамический блок штатными средствами AutoCAD.", True)
    step(doc, 1, "Распакуйте обновлённый ZIP в отдельную папку. Сохраните рабочие чертежи. Откройте полный AutoCAD 2024 и новый пустой несохранённый чертёж. Изменения только вида или окна (zoom/pan) допустимы; объекты и настройки чертежа перед запуском не меняйте.")
    step(doc, 2, 'Прочитайте настройку выражением (getvar "LISPSYS"): требуется 1 или 2. Оно не открывает запрос изменения. Если значение пришлось изменить с 0, перезапустите AutoCAD и снова откройте новый пустой DWG. Не запускайте автор в редакторе блока или в рабочем DWG.')
    step(doc, 3, "_.APPLOAD → выберите vector1_2015_4_2_1_training.lsp из новой папки. После загрузки в F2 должна быть редакция 2026-10-08.2. Защиту AutoCAD не отключайте.")
    step(doc, 4, "_.LOGFILEON → ATFNATIVEBUILD → новое имя DWG. Существующий файл не перезаписывается. Для русской версии AutoCAD вводите префикс _. в командах загрузки и журнала точно как здесь. Имя ATFNATIVEBUILD вводится без этого префикса.")
    step(doc, 5, "Дождитесь завершения. Автор создаёт параметры, действия, два состояния видимости и пять вставок. Затем проверяет геометрию, размеры, марки, изменение трёх вставок, две неизменные, повтор, возврат и копирование. DWG сохраняется только после этих проверок.")
    step(doc, 6, "После успешного завершения первые три вставки находятся в B, две оставшиеся — в A. Выполните _.ZOOM → _Extents, откройте F2 и сохраните текст результата.")
    paragraph(doc, "Если автор завершился ошибкой", True)
    paragraph(doc, "Видео 08.10 показывает отказ с UNDOCTL=61. Отдельный снимок — повтор при 53 и остановку на числе ручек Visibility. Ошибка Visibility и вторичный сбой обработчика относятся к автору, а не к действиям Германа.")
    paragraph(doc, "Сценарий с 61 пока не устранён: это признак открытой группы UNDO, владелец которой неизвестен. Автор не завершает её автоматически. Если отказ был до изменений, пустой DWG выбрасывать не нужно. Сохраните F2; не выполняйте ручной End ради обхода отказа и не обнуляйте DBMOD.")
    paragraph(doc, "Если построение началось и очистка не завершилась, закройте только тестовый DWG без сохранения. Сохраните полный F2: редакцию, UNDOCTL, DBMOD, DWGTITLED, BLOCKEDITOR, этап, команду и последний запрос. Геометрия после отказа не считается готовой библиотекой.")
    paragraph(doc, 'Путь к журналу: (getvar "LOGFILENAME"). После опыта выполните _.LOGFILEOFF, если включали журнал только для теста. Передайте Денису файл .log или полный F2.')

    page(doc, "2 Проверить три из пяти")
    paragraph(doc, "Проверяем реальные контуры, измеряемые расстояния и видимые марки. Одного изменения цифры в Properties недостаточно.", True)
    table(doc, ["Поле в Properties (Ctrl+1)", "A", "B"], [
        ("AFN_INSULATION", "100", "150"),
        ("AFN_CLADDING_X", "230", "280"),
        ("AFN_VARIANT", "V1_KR2_70_200_\nGP_40_40_1p2", "V1_KR2_70_250_\nGP_60_40_1p2"),
    ], [7.0, 5.2, 5.2])
    step(doc, 1, "Выберите все пять вставок и откройте Ctrl+1. Сначала задайте AFN_VARIANT=A, затем AFN_INSULATION=100 и AFN_CLADDING_X=230. Здесь A/B обозначают полные строки варианта из таблицы, а не буквальный ввод одной буквы.")
    step(doc, 2, "Снимите выбор. Сделайте снимок всех пяти исходных A. Рамкой выберите только три вставки; убедитесь, что выделены именно три блока.")
    step(doc, 3, "У выбранных трёх задайте сначала вариант B, затем утеплитель 150 и вынос 280. Геометрия утеплителя, плиты и дискретных деталей должна измениться согласованно; размеры — показать новые расстояния, марки — новый КР2/ГП.")
    step(doc, 4, "Проверьте две невыбранные: их контуры, размеры 100/230, КР2-70-200 и ГП-40-40-1,2 должны сохраниться. Базы вставок, повороты, масштабы и пользовательские атрибуты не должны сдвинуться.")
    step(doc, 5, "Повторите те же значения B: дополнительного смещения быть не должно. Верните выбранные три в A тем же порядком; сравните с исходным снимком.")
    step(doc, 6, "Отдельно проверьте линейные ручки, _.COPY и _.UNDO. После копирования изменение копии не должно менять исходник. Properties изменяет поля по отдельности: отменяйте нужное число операций до исходного состояния.")
    step(doc, 7, "Сохраните DWG, закройте его и откройте заново. Повторите переход трёх выбранных вставок в B. Сравните геометрию, марки и размеры до и после повторного открытия.")

    if downloads:
        page(doc, f"3 Скачать комплект {release_build}")
        paragraph(doc, f"Для команд библиотеки и проверки фасадных исправлений нужны AFacades + AClad + AFrame из одного выпуска №{release_build}. Если этот комплект уже установлен и работает, повторная установка ради обновлённого LISP не требуется. При обновлении модулей заменяются все три вместе.", True)
        paragraph(doc, "Обе общие инструкции DOCX вложены в тестовый ZIP и доступны без скачивания с GitHub. Ниже — ссылки на выпуск. Если модули ещё не установлены, откройте Install_AutoCAD_2024.docx, закройте AutoCAD и установите три бандла по инструкции.")
        downloads_table = table(doc, ["Скачать файл", "Назначение"], [(name, purpose) for name, purpose, _ in downloads], [7.3, 10.1])
        for row, (name, _, url) in zip(downloads_table.rows[1:], downloads):
            p = row.cells[0].paragraphs[0]
            p.clear()
            hyperlink(p, name, url)
        step(doc, 1, "Если требуется установка модулей, скачайте все три ZIP из таблицы. Source code не нужен; модули разных выпусков не смешивайте.")
        step(doc, 2, f"После установки запустите AutoCAD и проверьте в F2 номер {release_build} у всех трёх модулей. Затем переходите к командам следующего раздела.")
        paragraph(doc, "Тестовый ZIP содержит LISP-генератор и материалы опыта; установочных бандлов и готового динамического DWG в нём нет. Обновление DLL не подтверждает успешный запуск автора: нативный узел по-прежнему нужно создать и проверить в AutoCAD.")
    page(doc, ("4" if downloads else "3") + " Плагин и результаты проверки")
    if downloads:
        paragraph(doc, f"Команды ниже проверяйте после установки трёх модулей выпуска №{release_build} по ссылкам предыдущего раздела. Выпуск №108 этих команд не содержит.", True)
    else:
        paragraph(doc, "Команды ниже требуют исправленного комплекта AFacades + AClad + AFrame. Выпуск №108 их не содержит. Номер нового установочного выпуска здесь не указан; эта инструкция и LISP не обновляют установленные DLL.", True)
    table(doc, ["Команда", "Что проверить после обновления модулей"], [
        ("ATFNODEIMPORT", "Выбрать полученный библиотечный DWG, затем точку вставки. Чужой блок с совпавшим именем не должен быть заменён."),
        ("ATFNODEDEMO", "Выбрать одну библиотечную вставку; получить пять её копий для отдельного опыта."),
        ("ATFNODEEDIT", "Ввести утеплитель, вынос и A/B, затем выбрать вставки. Проверить три выбранных и две неизменные; одна команда — одна атомарная запись."),
        ("ATFNODETEST", "Пять одинаковых A → выбрать все пять → ввести 150 и 280. Команда меняет три первых по handles, проверяет результат и откатывает свой опыт."),
    ], [4.2, 13.2])
    paragraph(doc, "ATFNODE остаётся справочной схемой слоёв. Для библиотеки используются команды из таблицы. Автоматический тест не заменяет ручки, ручной COPY/Undo и сохранение/открытие.")
    doc.add_heading("Что передать Денису", 2)
    paragraph(doc, "Заполните Test_Results.txt. Приложите DWG, снимки «пять A» и «три B + две A», полный F2/журнал и точные действия до ошибки. При отказе укажите, на каком шаге он произошёл; отказ обычного сценария — ошибка опыта, а не успешный тест.")
    paragraph(doc, f"Исправления и порядок повторной проверки описаны в {fixes_name}. Отдельно отмечен неустранённый сценарий UNDOCTL=61. Синтетические проверки и Windows CI не являются выполнением динамического узла в AutoCAD.")
    paragraph(doc, "Не принимайте учебные значения за инженерное назначение: параметры проекта, крепёж, монтажные диапазоны и расчётное подтверждение задаются отдельно.")
    doc.add_heading("Идентификация комплекта", 2)
    paragraph(doc, "LISP-автор: 2026-10-08.2. Эта редакция должна быть видна в F2 после загрузки.")
    paragraph(doc, "Исходный код: " + code_sha)
    paragraph(doc, "Windows CI: " + ci_url)
    if downloads:
        paragraph(doc, f"Установочные модули и две инструкции: build-{release_build}. Ссылки закреплены за этим номером; latest не используется.")
    paragraph(doc, "Успешная нативная приёмка в AutoCAD: НЕ ЗАВЕРШЕНА. Итог обновлённого LISP фиксирует проектная группа после этого маршрута.", True)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--diagram", type=Path)
    parser.add_argument("--code-sha", required=True)
    parser.add_argument("--ci-url", required=True)
    parser.add_argument("--release-build", type=int,
                        help="Explicit published build 110 or newer; omit for the offline pilot without release links")
    parser.add_argument("--fixes-name", default="Fixes_2026-10-08_09_build111.txt",
                        help="Name of the fixes checklist included in the pilot ZIP")
    args = parser.parse_args()
    build(args.out, args.diagram, args.code_sha, args.ci_url, args.release_build, args.fixes_name)
    print(args.out)
