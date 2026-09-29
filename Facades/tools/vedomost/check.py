# -*- coding: utf-8 -*-
"""Проверка ядра ведомостей работ (29.09n): mono-прогон Check.cs → книги в /tmp/ved_out → openpyxl.
Запуск (из корня): python3 Facades/tools/vedomost/check.py  (нужны mono/mcs и openpyxl)."""
import os
import subprocess
import sys
import zipfile
import openpyxl

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", ".."))
TPL = os.path.join(ROOT, "Facades", "bundle", "AFacades.bundle", "Contents", "templates")
OUT = "/tmp/ved_out"
fails = []


def ok(cond, msg):
    print(("OK   " if cond else "FAIL ") + msg)
    if not cond:
        fails.append(msg)


def main():
    exe = "/tmp/vedomost.exe"
    subprocess.check_call(["mcs", "-out:" + exe, "-r:System.Web.Extensions.dll", "-r:System.IO.Compression.dll",
                           "-r:System.Xml.dll", os.path.join(ROOT, "Facades", "tools", "vedomost", "Check.cs"),
                           os.path.join(ROOT, "Facades", "src", "AFacadesPlugin", "WorkStatement.cs")],
                          stdout=subprocess.DEVNULL)
    subprocess.check_call(["mono", exe, TPL, OUT], stdout=subprocess.DEVNULL,
                          env=dict(os.environ, LANG="C.UTF-8"))
    D = lambda ws, r: ws["D%d" % r].value
    # 1. штукатурка: D3/D9/D10/D12, D11 пусто + «заполнить вручную», утеплитель с толщиной и =D3
    ws = openpyxl.load_workbook(os.path.join(OUT, "sht_one.xlsx")).worksheets[0]
    ok((D(ws, 3), D(ws, 9), D(ws, 10), D(ws, 12)) == (1000, 250, 56, 45.2), "штукатурка: D3/D9/D10/D12 = 1000/250/56/45.2")
    ok(D(ws, 11) is None and ws["E11"].value == "заполнить вручную", "штукатурка: армирование углов — пусто, «заполнить вручную»")
    ok("толщиной 150 мм на клей" in ws["B4"].value and D(ws, 4) == "=D3", "штукатурка: утеплитель 150 мм, формула =D3 осталась")
    ok([D(ws, r) for r in range(5, 9)] == ["=D3"] * 4 and ws["A12"].value == "1.10", "штукатурка: формулы 5–8 и позиции не тронуты")
    # 2. вент, два утеплителя: строка 5 (150) и 6 (100), ниже всё на строку
    ws = openpyxl.load_workbook(os.path.join(OUT, "vent_two.xlsx")).worksheets[0]
    ok(D(ws, 3) == 1250 and (D(ws, 5), D(ws, 6)) == (1000, 250), "вент: D3 1250, утеплители 1000 (150) и 250 (100)")
    ok("150 мм" in ws["B5"].value and "100 мм" in ws["B6"].value and ws["A5"].value == "2.1" and
       ws["A6"].value == "2.2" and ws["A7"].value == "2.3" and "мембраны" in ws["B7"].value, "вент: позиции 2.1/2.2/2.3")
    ok(D(ws, 7) == "=D3" and D(ws, 9) == "=D3" and D(ws, 10) == "=D3", "вент: мембрана и НВФ по-прежнему =D3")
    ok((D(ws, 12), D(ws, 14), D(ws, 17), D(ws, 21)) == (62, 280, 42, 30), "вент: отливы/откосы/витражи/парапет на сдвинутых строках")
    ok((D(ws, 13), D(ws, 15), D(ws, 16), D(ws, 18), D(ws, 19)) == ("=D12", "=D14", "=D14", "=D17", "=D17"),
       "вент: формулы клипс и отсечек сдвинуты вместе со строками")
    ok(D(ws, 20) is None and ws["E20"].value == "заполнить вручную" and "цокольного" in ws["B20"].value,
       "вент: цокольный отлив — вручную")
    ok(ws["A12"].value == "4.1" and ws["A21"].value == "4.10", "вент: позиции раздела 4 не тронуты")
    # 3. повтор по той же книге: строки 150/100 уже есть — копий нет, площади на своих местах
    ws = openpyxl.load_workbook(os.path.join(OUT, "vent_again.xlsx")).worksheets[0]
    ok((D(ws, 5), D(ws, 6)) == (1000, 250) and ws.max_row == openpyxl.load_workbook(os.path.join(OUT, "vent_two.xlsx")).worksheets[0].max_row,
       "повтор: строки с толщиной найдены по тексту, новых строк нет")
    # 4. зона старой сборки: откосы/отливы — как у окон
    ws = openpyxl.load_workbook(os.path.join(OUT, "vent_old.xlsx")).worksheets[0]
    ok((D(ws, 3), D(ws, 11), D(ws, 13), D(ws, 16)) == (100, 12, 40, 0), "старая зона: все проёмы — окна")
    # книга: пересчёт при открытии, нет calcChain и внешних ссылок
    for f in ("sht_one.xlsx", "vent_two.xlsx"):
        z = zipfile.ZipFile(os.path.join(OUT, f))
        wbx = z.read("xl/workbook.xml").decode("utf-8")
        ok('fullCalcOnLoad="1"' in wbx and "xl/calcChain.xml" not in z.namelist() and
           not any("external" in n for n in z.namelist()), "%s: пересчёт при открытии, без calcChain и внешних ссылок" % f)
    print("ВЕДОМОСТИ: " + ("OK" if not fails else "провалов %d" % len(fails)))
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
