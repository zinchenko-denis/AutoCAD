"""Current placement guard and compatibility of archived unlinked-bracket reports.

The current engine must reject null-step placement with fewer than two supports.
The archived response and explicit DTO variants exercise the actual producer and
table for old drawings; they do not represent accepted current placement.
Synthetic dimensions and the matching tolerance are not mounting design limits.
"""
import argparse
import copy
import hashlib
import json
from pathlib import Path
import subprocess
import sys

ROOT = Path(__file__).resolve().parents[2]
for name in ('AFrame/engine', 'tools/catalog_connections_0110', 'tools/quantities_3009'):
    sys.path.insert(0, str(ROOT / name))
import frame_engine
from frame_connections import build_connection_passport, merge_connection_passports
from test_frame_solution_selection import example_selection
from test_connection_store import compile_actual_producer
from probe_runtime import command, compile_probe, available


LEGACY_INPUT = ROOT / 'docs/orphan_connections_0110/reproduction/fixed_minimal_input.json'


def legacy_variant(fixture, req, gap, zones):
    """Stored old result, or explicit geometry variants; never a new engine result."""
    response = copy.deepcopy(fixture['response'])
    if gap == 10 and zones == 1:
        return response  # Exact archived response, including its stored passport.
    template = copy.deepcopy(response)
    response['rails'] = []; response['brackets'] = []; response['per_zone'] = []
    response['design_scope']['per_zone'] = []
    passports = []
    for zone in req['zones']:
        zone_id = zone['zone_id']
        local = copy.deepcopy(template)
        local['rails'][0]['y1'] = round(1500 - gap / 2, 4)
        local['rails'][1]['y0'] = round(1500 + gap / 2, 4)
        for rail in local['rails']:
            rail['len'] = round(rail['y1'] - rail['y0'], 4)
            rail['zone'] = zone_id
        local['brackets'][0]['zone'] = zone_id
        passports.append((zone_id, build_connection_passport(req, local),
                          len(response['rails']), len(response['brackets'])))
        response['rails'].extend(local['rails']); response['brackets'].extend(local['brackets'])
        per_zone = copy.deepcopy(template['per_zone'][0])
        per_zone.update(zone_id=zone_id, rails_lm=round((3000 - gap) / 1000, 6))
        response['per_zone'].append(per_zone)
        scope = copy.deepcopy(template['design_scope']['per_zone'][0]); scope['zone_id'] = zone_id
        response['design_scope']['per_zone'].append(scope)
    response['system_used']['rail_gap'] = gap
    response['summary'].update(rails=2 * zones, rails_lm=round((3000 - gap) * zones / 1000, 6),
                               brackets_main=zones, zones=zones)
    response['connection_passport'] = merge_connection_passports(req, passports)
    return response


def check_null_step_refusal(response, gap):
    assert response.get('ok') is False and response.get('error_code') == 'E_UNSUPPORTED_RAIL', response
    assert response['failed_zone'] == 'Z0', response
    assert all(not response.get(key) for key in ('rails', 'hrails', 'brackets', 'clamps', 'fittings')), response
    expected_supports = [1500.0] if gap <= 1 else []
    expected = [(0, round(1500 - gap / 2, 4)), (round(1500 + gap / 2, 4), 3000)]
    assert response['unsupported_counts'] == {'rail_pieces': 2}, response
    assert len(response['unsupported']) == 2, response
    for index, (member, (y0, y1)) in enumerate(zip(response['unsupported'], expected)):
        assert member == dict(member_index=index, x=600.0, y0=y0, y1=y1,
                              length=round(y1 - y0, 4), support_y=expected_supports,
                              support_count=len(expected_supports), bracket_start_offset=300.0), member


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out', type=Path, required=True)
    parser.add_argument('--expect', choices=('fixed',), default='fixed',
        help='Before-state replay uses docs/orphan_connections_0110/reproduction/baseline_reproduce.py against the recorded source checkout')
    args = parser.parse_args()
    if not available():
        print('BLOCKED: native .NET compiler/runtime unavailable')
        return 2
    out = args.out.resolve(); out.mkdir(parents=True, exist_ok=True)
    legacy_input = json.loads(LEGACY_INPUT.read_text(encoding='utf-8'))
    base = {'op': 'frame', 'system': {'name': 'Вектор-1', 'rail_std': 1500, 'rail_gap': 10,
        'bracket_step': None, 'bracket_step_corner': None}, 'sub_type': 'vertical',
        'cladding': 'porcelain', 'parts': 'frame', 'solution_selection': example_selection(),
        'floors_y': [1500], 'joints_x': [600], 'zones': [{'zone_id': 'Z0', 'zone': {
            'contour': {'pts': [[0, 0], [1200, 0], [1200, 3000], [0, 3000]]}, 'openings': []}}]}
    producer, producer_sources = compile_actual_producer(out / 'producer')
    inputs = {}
    for name, gap, steps, zones in (('minimal', 10, None, 1), ('two_same_coordinate_zones', 10, None, 2),
            ('gap_zero', 0, None, 1), ('gap_one', 1, None, 1), ('gap_one_plus', 1.0002, None, 1),
            ('normal_positive_steps', 10, 800, 1)):
        req = copy.deepcopy(base)
        req['system'].update(rail_gap=gap, bracket_step=steps, bracket_step_corner=steps)
        if zones == 2:
            req['zones'].append(copy.deepcopy(req['zones'][0])); req['zones'][1]['zone_id'] = 'Z1'
        directory = out / name; directory.mkdir(exist_ok=True)
        current_response = frame_engine.run(copy.deepcopy(req))
        if steps is None:
            check_null_step_refusal(current_response, gap)
            (directory / 'current_refusal.json').write_text(json.dumps(current_response, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
            response = legacy_variant(legacy_input, req, gap, zones)
            origin = 'archived_legacy_response' if name == 'minimal' else 'explicit_archived_response_variant'
            if name == 'minimal':
                assert req == legacy_input['request'] and response == legacy_input['response']
        else:
            response = current_response
            assert response['ok'], response
            assert all(m['support_count'] >= 2 for m in response['connection_passport']['members'])
            origin = 'current_engine_response'
        source = directory / 'input.json'
        source.write_text(json.dumps({'request': req, 'response': response, 'mapping': {}}, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        produced_path = directory / 'producer_output.json'
        subprocess.run(command(producer) + ['--produce', str(source), str(produced_path)], check=True)
        produced = json.loads(produced_path.read_text(encoding='utf-8'))
        assert produced['ok'], produced
        report_path = directory / 'report.json'
        report_path.write_text(json.dumps(produced['report'], ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        linked = {i for member in response['connection_passport']['members'] for support in member['supports'] for i in support['bracket_indices']}
        orphan_indices = sorted(set(range(len(response['brackets']))) - linked)
        expected = [] if name in ('gap_zero', 'gap_one', 'normal_positive_steps') else list(range(zones))
        assert orphan_indices == expected, (name, orphan_indices)
        inputs[name] = {'origin': origin, 'current_engine_status': 'E_UNSUPPORTED_RAIL' if steps is None else 'ok',
            'gap': gap, 'step': steps, 'zones': zones, 'rail_count': len(response['rails']),
            'bracket_count': len(response['brackets']), 'orphan_indices': orphan_indices,
            'report': str(report_path), 'input_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
            'report_sha256': hashlib.sha256(report_path.read_bytes()).hexdigest()}
    sources = [Path(__file__).with_name('OrphanTableProbe.cs'), ROOT / 'Common/FacadeQuantities.cs']
    sources += [ROOT / 'Facades/src/AFacadesPlugin' / name for name in ('ConnectionTableData.cs',
        'QuantityTableIdentity.cs', 'QuantityTableView.cs', 'QuantityTableView.Connections.cs',
        'CladdingTableData.cs', 'FrameTableData.cs', 'QuantityXlsxFile.cs', 'XlsxWriter.cs')]
    exe = out / 'OrphanTableProbe.exe'
    compiled = compile_probe(sources, ['System.Web.Extensions', 'System.Xml', 'System.IO.Compression', 'System.IO.Compression.FileSystem'], exe)
    (out / 'compile.log').write_text(compiled.stdout + compiled.stderr, encoding='utf-8')
    if compiled.returncode:
        print(compiled.stdout + compiled.stderr)
        return 1
    checks = 0; cases = []
    for name in inputs:
        mode = 'scale' if name == 'minimal' else 'full' if name == 'two_same_coordinate_zones' else 'controls'
        run = subprocess.run(command(exe) + [inputs[name]['report'], str(out / name / 'scope'), args.expect, mode], text=True, capture_output=True)
        (out / name / 'scope.log').write_text(run.stdout + run.stderr, encoding='utf-8')
        assert run.returncode == 0, run.stdout + run.stderr
        observed = json.loads((out / name / 'scope/cases.json').read_text(encoding='utf-8'))
        checks += observed['checks']
        cases.append({'fixture': name, **observed})
    from openpyxl import load_workbook
    exports = []
    for path in sorted(out.glob('*/scope/*.xlsx')):
        if path.name.endswith('_roundtrip.xlsx'):
            continue
        book = load_workbook(path)
        values = {sheet.title: list(sheet.values) for sheet in book}
        maximum = max((len(cell.value) for sheet in book for row in sheet for cell in row if isinstance(cell.value, str)), default=0)
        assert maximum <= 32767, (path, maximum)
        roundtrip = path.with_name(path.stem + '_roundtrip.xlsx')
        book.save(roundtrip)
        restored = load_workbook(roundtrip)
        assert {sheet.title: list(sheet.values) for sheet in restored} == values, path
        exports.append({'path': str(path.relative_to(out)), 'maximum_cell_characters': maximum,
            'roundtrip_exact': True, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})
    dependencies = set(producer_sources + sources + [Path(__file__).resolve(), LEGACY_INPUT, ROOT / 'tools/catalog_connections_0110/test_connection_store.py',
        ROOT / 'tools/quantities_3009/probe_runtime.py', ROOT / 'AFrame/engine/systems.json'])
    dependencies.update((ROOT / 'AFrame/engine').glob('*.py'))
    manifest = {'status': 'PASS' if args.expect == 'fixed' else 'REPRODUCED', 'expected': args.expect,
        'inputs': inputs, 'checks': checks, 'cases': cases, 'xlsx_exports': exports,
        'current_null_step_refusals': 5,
        'legacy_fixture': str(LEGACY_INPUT.relative_to(ROOT)),
        'path': 'archived response / explicit compatibility variants / current positive-step engine -> actual FrameQuantities producer -> common validator -> table -> XLSX',
        'scope_variants': 'gap/zone variants modify the archived DTO before the producer and rebuild inventory-only passports; retained/legacy/unavailable/alias/removed-element variants modify typed DTOs after the producer',
        'current_engine_responses_modified': False,
        'legacy_null_step_input': 'rejected by current engine; archived report compatibility only',
        'engineering_status': 'unchanged; mounting not confirmed', 'live_autocad_checked': False,
        'source_sha256': {str(path.relative_to(ROOT)): hashlib.sha256(path.read_bytes()).hexdigest() for path in sorted(dependencies)}}
    (out / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print('Current null-step refusals: 5 PASS; orphan compatibility checks:', checks, 'PASS; XLSX exports/roundtrips:', len(exports))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
