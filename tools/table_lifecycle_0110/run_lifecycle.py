"""Real table command branch on the limited pilot's same persistent engine/store DB.

CAD/prompt/form objects are explicit doubles. No native DWG, UI or transaction
rollback is claimed. --expect baseline verifies the known UCS defect only.
"""
import argparse
import hashlib
import importlib.util
import json
from pathlib import Path
import sys
import time

ROOT=Path(__file__).resolve().parents[2]
spec=importlib.util.spec_from_file_location('limited_pilot',ROOT/'tools/limited_pilot_0110/run_acceptance.py')
pilot=importlib.util.module_from_spec(spec);spec.loader.exec_module(pilot)


def helper_extraction(build):
    source=ROOT/'Facades/src/AFacadesPlugin/ZoneCommand.cs';raw=source.read_text(encoding='utf-8')
    start=raw.index('        internal static void StoreZoneData(Transaction tr, Entity ent,\n                                           string json, string key, IEnumerable<ObjectId> refs)')
    end=raw.index('\n        private static void RemoveZoneData',start)
    store=raw[start:end]
    start=raw.index('        internal static string ReadZoneData(Transaction tr, Entity ent,\n                                            string key)')
    end=raw.index('\n        // ── служебное',start)
    read=raw[start:end]
    start=raw.index('        internal static string SafeStr(object o)');end=raw.index('\n    }',start)
    safe=raw[start:end]
    output=build/'ZoneMetadataHelpers.cs'
    output.write_text('using System;using System.Collections.Generic;using System.Globalization;using System.Text;using Autodesk.AutoCAD.DatabaseServices;\nnamespace AFacadesPlugin {internal static class ZoneCommand {\n'+store+'\n'+read+'\n'+safe+'\n}}\n',encoding='utf-8')
    return source,output


def verify_tables(out):
    from openpyxl import load_workbook
    checks=[]
    def need(ok,msg):
        if not ok:raise AssertionError(msg)
        checks.append(msg)
    results=json.loads((out/'table_command/result.json').read_text(encoding='utf-8'))
    oracle=json.loads((ROOT/'docs/table_lifecycle_0110/independent_oracle.json').read_text(encoding='utf-8'))
    for position in oracle['position_cases']:
        name='ucs_reproduction_frame' if position['name']=='translated_rotated' else 'position_'+position['name']
        observed=next(c for c in results['cases'] if c['name']==name)
        expected=position['picked_ucs'] if results['expect']=='baseline' else position['expected_wcs']
        need(all(abs(a-b)<1e-10 for a,b in zip(observed['actual_position'],expected)), 'independent coordinate oracle: '+position['name'])
    def cell_text(value):
        # Fixture-level independent formatting oracle: all its numeric values
        # are finite ordinary quantities, with six decimal places for display.
        if value is None:text=''
        elif isinstance(value,float):text=format(value,'.6f').rstrip('0').rstrip('.') or '0'
        else:text=str(value)
        return text.replace('\r\n','\n').replace('\r','\n').replace('\\','\\\\').replace('{','\\{').replace('}','\\}').replace('%','\\U+0025').replace('\n','\\P')
    files=[]
    for path in sorted((out/'table_command').glob('*_rows.json')):
        stem=path.name.removesuffix('_rows.json');expected=json.loads(path.read_text(encoding='utf-8'))
        width=max(map(len,expected))
        table_path=path.parent/(stem+'_table.json')
        if table_path.exists():
            table=json.loads(table_path.read_text(encoding='utf-8'))
            need(table['metadata']['rows']==expected,'actual stored table metadata equals preview rows: '+stem)
            need(table['row_count']==len(expected) and table['column_count']==width,'actual Table dimensions match complete expected rectangle: '+stem)
            body=json.loads(table['body'])['rows']
            need(len(body)==len(expected) and all(len(r)==width for r in body),'actual Table cells have no extra rows/columns: '+stem)
            for r,row in enumerate(expected):
                for c in range(width):need(body[r][c]==cell_text(row[c] if c<len(row) else None),'actual Table cell equals independently formatted/escaped source '+stem+':'+str(r)+':'+str(c))
            need(table['merges']==[[r,0,r,width-1] for r,row in enumerate(expected) if len(row)==1],'actual Table merged rows match single-cell headings/notes: '+stem)
            need(any('\\{A\\}\\\\путь 50\\U+0025\\P' in cell for row in body for cell in row),'actual Table preserves escaped braces/backslash/percent/newline in user note: '+stem)
        book=load_workbook(path.parent/(stem+'.xlsx'));sheet=book.active
        need(sheet.max_row==len(expected) and sheet.max_column==width,'XLSX dimensions exactly match complete expected rectangle: '+stem)
        for r,row in enumerate(expected,1):
            for c in range(1,width+1):
                value=row[c-1] if c<=len(row) else None
                need(sheet.cell(r,c).value==value and sheet.cell(r,c).data_type!='f','command XLSX equals exact preview value without extra values/formulas '+stem+':'+str(r)+':'+str(c))
        need(all(not isinstance(cell.value,str) or len(cell.value)<=32767 for row in sheet for cell in row),'command XLSX cells fit Excel limits: '+stem)
        files.append({'name':stem,'xlsx_sha256':pilot.digest(path.parent/(stem+'.xlsx')),'rows':sheet.max_row,'columns':sheet.max_column})
        book.close()
    need(len(files)==10,'three real command output types at all three issued epochs plus Excel-only control')
    pilot.save(out/'table_command/xlsx_review.json',{'status':'PASS','checks':len(checks),'files':files,'assertions':checks})
    return results,len(checks)


def main():
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--out',required=True,type=Path)
    ap.add_argument('--expect',choices=['baseline','fixed'],default='fixed')
    ap.add_argument('--command-source',type=Path,help='Exact saved historical FacadeQuantityTableCommand.cs for baseline replay; no checkout/reset')
    args=ap.parse_args();out=args.out.resolve();out.mkdir(parents=True,exist_ok=True);build=out/'build';build.mkdir(exist_ok=True)
    if not pilot.available():print('BLOCKED: native .NET compiler/runtime unavailable');return 2
    command_source=(args.command_source or ROOT/'Facades/src/AFacadesPlugin/FacadeQuantityTableCommand.cs').resolve()
    command_hash=pilot.digest(command_source)
    baseline_commit='16c7307ce42715721ee7835d2dc90955e2229781'
    baseline_path='Facades/src/AFacadesPlugin/FacadeQuantityTableCommand.cs'
    baseline_lf_sha256='198cb83d4fb7c99cc24da4fba81c6cf4c6c14d20d34b320b2da57b72f4fde922'
    normalized_command_hash=hashlib.sha256(command_source.read_bytes().replace(b'\r\n',b'\n')).hexdigest()
    baseline_matches=normalized_command_hash==baseline_lf_sha256
    if args.expect=='baseline' and not baseline_matches:
        raise AssertionError('Baseline reproduction requires the exact reviewed 16c7307 command source (LF/CRLF only)')
    # Preserve the exact full command used by this run, including before-fix replay.
    (out/'command_source.cs').write_bytes(command_source.read_bytes())
    original,generated=helper_extraction(build)
    extra=[Path(__file__).with_name('TableCommandDoubles.cs'),Path(__file__).with_name('TableLifecycleProbe.cs'),command_source]
    tick=time.perf_counter();exe,sources,extracted=pilot.compile_native(build,extra_sources=extra,extra_generated=[generated]);compile_seconds=time.perf_counter()-tick
    tracked=[p for p in sources if p.is_relative_to(ROOT)]+[original,Path(__file__).resolve(),ROOT/'tools/limited_pilot_0110/run_acceptance.py',ROOT/'tools/quantities_3009/probe_runtime.py',ROOT/'docs/limited_pilot_0110/independent_oracle.json',ROOT/'docs/table_lifecycle_0110/independent_oracle.json']
    for directory in ('AFrame/engine','AClad/engine','Facades/engine'):
        tracked.extend((ROOT/directory).glob('*.py'));tracked.extend((ROOT/directory).glob('*.json'))
    hashes={str(p.relative_to(ROOT)):pilot.digest(p) for p in sorted(set(tracked))}
    selection=pilot.example_selection();selection['geometry'].update(cladding_front_offset_mm=230,insulation_layers_mm=[100,50])
    pilot.save(out/'selection.json',selection);pilot.save(out/'table_expect.json',{'expect':args.expect})
    tick=time.perf_counter();observations=pilot.execute(exe,out/'selection.json',out);duration=time.perf_counter()-tick
    pilot.verify(out,observations);pilot.verify_trace(out)
    table,table_python=verify_tables(out)
    native=json.loads((out/'native_result.json').read_text(encoding='utf-8'))
    unchanged=all(pilot.digest(ROOT/name)==value for name,value in hashes.items()) and pilot.digest(command_source)==command_hash
    pilot.need(unchanged,'all compiled/source/extraction inputs remain unchanged')
    manifest={'schema':'table_lifecycle_acceptance/1','status':'PASS','expect':args.expect,
        'verification':'PASS' if args.expect=='fixed' else 'BASELINE_DEFECT_REPRODUCED',
        'engineering_assessment':'BLOCKED','full_P5':'NOT_RUN','live_autocad_checked':False,
        'base_native_checks':native['checks'],'base_python_checks':len(pilot.CHECKS),
        'table_native_checks':table['checks'],'table_python_checks':table_python,
        'checks':native['checks']+len(pilot.CHECKS)+table['checks']+table_python,
        'engine_invocations':len(observations),'command_xlsx_files':10,'base_xlsx_files':9,
        'source_sha256':hashes,'actual_command_sha256':command_hash,
        'command_source_origin':{'mode':'explicit_source_override' if args.command_source else 'working_repository_source','saved_copy':'command_source.cs','matches_baseline_commit':baseline_commit if baseline_matches else None,'baseline_path':baseline_path,'baseline_git_blob_lf_sha256':baseline_lf_sha256,'actual_command_lf_sha256':normalized_command_hash},
        'generated_extraction_sha256':{str(p.relative_to(out)):pilot.digest(p) for p in extracted},
        'sources_unchanged_during_run':unchanged,
        'artifact_sha256':{str(p.relative_to(out)):pilot.digest(p) for p in sorted(out.rglob('*')) if p.is_file() and build not in p.parents and p.suffix in ('.json','.xlsx','.cs') and p.name!='manifest.json'},
        'timings':{'compile_seconds':compile_seconds,'persistent_pipeline_seconds':duration,'scope':'managed adapters only'},
        'limits':['actual FacadeQuantityTableCommand.Run branch; external ATFTABLE dispatcher and native WinForms NOT_RUN',
            'actual helpers/producers/guards/stores in same persistent DB, scripted prompts and Table model',
            'QuantityCadDoubles transactions have no rollback; cancellation tested only before table writes',
            'native DWG rendering/save/open/COPY/Undo and full ATFNODE lifecycle NOT_RUN',
            'final late-project control deliberately commits revision3; it is not rolled back or relabelled as the earlier issued epoch']}
    pilot.save(out/'manifest.json',manifest)
    print('Table lifecycle:',manifest['checks'],manifest['verification'],'; engineering BLOCKED; native host NOT_RUN')
    return 0


if __name__=='__main__':raise SystemExit(main())
