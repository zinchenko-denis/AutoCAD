"""Execute connection presentation and exact XLSX export without AutoCAD.

Checks physical identities, missing/unsupported states and update discrimination,
then generates a fresh engine -> actual producer report for end-to-end export.
"""
import argparse
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'tools/quantities_3009'))
from probe_runtime import available, command, compile_probe


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out', type=Path)
    parser.add_argument('--report', type=Path)
    args = parser.parse_args()
    if not available():
        print('BLOCKED: native .NET compiler/runtime unavailable; no table execution claimed')
        return 2
    out = (args.out or Path(tempfile.mkdtemp(prefix='connection_table_'))).resolve()
    out.mkdir(parents=True, exist_ok=True)
    sources = [Path(__file__).with_name('ConnectionTableProbe.cs'), ROOT / 'Common/FacadeQuantities.cs']
    sources += [ROOT / 'Facades/src/AFacadesPlugin' / name for name in (
        'ConnectionTableData.cs', 'QuantityTableIdentity.cs', 'QuantityTableView.cs', 'QuantityTableView.Connections.cs',
        'CladdingTableData.cs', 'FrameTableData.cs', 'QuantityXlsxFile.cs', 'XlsxWriter.cs')]
    exe = out / 'ConnectionTableProbe.exe'
    compiled = compile_probe(sources, ['System.Web.Extensions', 'System.Xml', 'System.IO.Compression', 'System.IO.Compression.FileSystem'], exe)
    (out / 'compile.log').write_text(compiled.stdout + compiled.stderr, encoding='utf-8')
    if compiled.returncode:
        print(compiled.stdout + compiled.stderr)
        return 1
    subprocess.run(command(exe) + [str(out)], check=True)
    from openpyxl import load_workbook
    sheet = load_workbook(out / 'connection_exact.xlsx')['Соединения']
    rows = list(sheet.values)
    checks = 0

    def check(value, message):
        nonlocal checks
        checks += 1
        assert value, message

    first = next(row for row in rows if row[0] == 'А')
    check(first[2:5] == (3, 2, '1000.1234567; 1000.1234567'), 'Position/span counts or precise intervals changed')
    check(first[5:7] == (123.4567891234567, 234.5678912345678), 'Free length precision changed')
    check('physical-1:rail' in first[10] and 'CAD A9' in first[1], 'Physical identity or CAD handle missing')
    check('physical-1:bracket:0, physical-1:coincident:0' in first[10], 'Coincident physical links lost')
    check('Справочное совпадение' in first[9] and 'изделие не подобрано' in first[9], 'Reference became a verified SKU')
    check('перекрытие проекций 12.34567891234567' in first[8], 'Adjacency gap precision or sign changed')
    for zone in ('Б', 'В'):
        missing = next(row for row in rows if row[0] == zone)
        check(missing[2:4] == ('не определено', 'не определено'), 'Unavailable/legacy passport became zero')
    no_support = next(row for row in rows if row[0] == 'Г')
    check(no_support[2] == 0 and no_support[5:7] == ('Опор-кандидатов нет', 'Опор-кандидатов нет'), 'Zero support free-length convention misleading')
    check(rows[-1][0] == 'Примечание пользователя:   Проверено & <узел> =1+1  ', 'User note, XML or whitespace corrupted')
    check(all(cell.data_type != 'f' for row in sheet for cell in row), 'Text became an Excel formula')
    from test_connection_store import compile_and_run_fixture
    report_path = args.report or compile_and_run_fixture(out / 'producer_fixture')
    report = json.loads(report_path.read_text(encoding='utf-8'))
    subprocess.run(command(exe) + ['--report', str(report_path), str(out / 'connection_engine.xlsx')], check=True)
    real_sheet = load_workbook(out / 'connection_engine.xlsx')['Соединения']
    actual = list(real_sheet.values)
    real_checks = 0
    for member in (member for passport in report['connection_passports'] for member in passport['members']):
        row = next(row for row in actual if len(row) > 10 and str(row[10]).startswith('Элемент: ' + member['rail_element_id'] + '. '))
        check(row[2:4] == (member['support_count'], member['span_count']), 'Real engine counts changed by export')
        real_checks += 1
    check(all(cell.data_type != 'f' for row in real_sheet for cell in row), 'Real source text became an Excel formula')
    # Actual selected-pilot engine output exposes shared physical candidate IDs.
    # The existing producer and common validator are used without patching its response.
    sys.path.insert(0, str(ROOT / 'tools/mounting_contract_0110'))
    from reproduce_gap import make_fixture, setup
    setup(ROOT)
    shared_report_path = make_fixture(out / 'shared_support')
    subprocess.run(command(exe) + ['--shared-report', str(shared_report_path), str(out / 'shared_support')], check=True)
    shared_cases = json.loads((out / 'shared_support/shared_cases.json').read_text(encoding='utf-8'))
    shared_sheet = load_workbook(out / 'shared_support/shared_support.xlsx')['Соединения']
    shared_rows = list(shared_sheet.values)
    check(len([row for row in shared_rows if isinstance(row[0], str) and
        row[0].startswith('Монтажная принадлежность не определена: кронштейн ')]) == 4,
        'Shared physical candidate warnings absent or duplicated in XLSX')
    check(len([row for row in shared_rows if len(row) > 10 and isinstance(row[10], str) and
        'Монтажная принадлежность не определена: общих кандидатов-кронштейнов — ' in row[10]]) == 4,
        'Affected real member rows not addressed in XLSX')
    check(all(cell.data_type != 'f' for row in shared_sheet for cell in row), 'Shared warning text became an Excel formula')
    for case in shared_cases['cases']:
        if case['name'].startswith('synthetic_single_bracket_fanout_'):
            check(case['members_indexed'] == case['members'] and case['links_indexed'] == case['members'],
                'Shared candidate pass revisits member/link pairs')
            check(case['messages'] == 0 if case['members'] == 1 else case['messages'] >= 1,
                'Fanout loses shared peer list or emits false warning')
            check(case['max_shared_message_characters'] <= 8000, 'Shared warning exceeds output chunk size')
            fanout_path = out / 'shared_support' / ('fanout_%d.xlsx' % case['members'])
            fanout_book = load_workbook(fanout_path)
            original_values = {sheet.title: list(sheet.values) for sheet in fanout_book}
            check(all(not isinstance(cell.value, str) or len(cell.value) <= 32767
                for sheet in fanout_book for row in sheet for cell in row), 'Exported XLSX cell exceeds Excel text limit')
            roundtrip_path = fanout_path.with_name(fanout_path.stem + '_roundtrip.xlsx')
            fanout_book.save(roundtrip_path)
            roundtrip_book = load_workbook(roundtrip_path)
            check({sheet.title: list(sheet.values) for sheet in roundtrip_book} == original_values,
                'Independent openpyxl roundtrip truncated/changed fanout output')
            check(case['warning_characters'] < 650 * case['members'] + 500,
                'Shared peer IDs multiply output quadratically')
    for name in ('many_brackets_320', 'many_brackets_near_limit'):
        many_path = out / 'shared_support' / (name + '.xlsx')
        many_book = load_workbook(many_path)
        original_values = {sheet.title: list(sheet.values) for sheet in many_book}
        check(all(not isinstance(cell.value, str) or len(cell.value) <= 32767
            for sheet in many_book for row in sheet for cell in row),
            'Member note exceeds Excel limit for many shared bracket IDs: ' + name)
        if name == 'many_brackets_near_limit':
            members = [row for row in original_values['Соединения'] if len(row) > 10 and
                isinstance(row[10], str) and row[10].startswith('Элемент: ')]
            check(len(members) == 2 and all('Монтажная принадлежность не определена:' in row[7]
                and row[7].startswith('Не определено: неподвижное / подвижное')
                and 'Монтажная принадлежность не определена:' not in row[10] for row in members),
                'Near-limit source note was modified or addressed warning lost from constraint cell')
        many_roundtrip = many_path.with_name(many_path.stem + '_roundtrip.xlsx')
        many_book.save(many_roundtrip)
        check({sheet.title: list(sheet.values) for sheet in load_workbook(many_roundtrip)} == original_values,
            'Independent roundtrip truncated many-bracket IDs: ' + name)
    sources += [ROOT / 'tools/mounting_contract_0110/reproduce_gap.py']
    sources += [ROOT / 'AFrame/src/AFramePlugin/FrameQuantities.cs', Path(__file__).with_name('test_connection_store.py'),
                Path(__file__).with_name('ConnectionStoreProbe.cs'), ROOT / 'tools/frame_quantities_3009/engine_oracles.py']
    sources += list((ROOT / 'AFrame/engine').glob('*.py')) + [ROOT / 'AFrame/engine/systems.json']
    manifest = {'status': 'PASS', 'xlsx_checks': checks, 'real_members_checked': real_checks,
                'shared_candidate_checks': shared_cases['checks'], 'shared_candidate_cases': shared_cases['cases'],
                'shared_fixture_sha256': hashlib.sha256(shared_report_path.read_bytes()).hexdigest(),
                'report_origin': 'supplied_report' if args.report else 'fresh_real_engine_to_actual_producer',
                'report_sha256': hashlib.sha256(report_path.read_bytes()).hexdigest(),
                'live_autocad_checked': False,
                'source_sha256': {str(path.relative_to(ROOT)): hashlib.sha256(path.read_bytes()).hexdigest() for path in sources + [Path(__file__).resolve()]}}
    (out / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print('Connection table XLSX checks:', checks, 'PASS; real engine members:', real_checks)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
