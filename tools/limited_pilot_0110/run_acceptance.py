"""One persistent managed CAD-double acceptance path, not a live AutoCAD session.

Actual settings/stores/producers and Python engines share source identities.
PASS verifies the bounded adapter path; mounting/3A stay BLOCKED and complete
P5 (native node command/store, DWG save/reopen/render/Undo/COPY) is NOT_RUN.
"""
import argparse
import copy
import hashlib
import json
from pathlib import Path
import queue
import re
import subprocess
import sys
import threading
import time

ROOT = Path(__file__).resolve().parents[2]
for directory in ('Facades/engine', 'AClad/engine', 'AFrame/engine', 'tools/quantities_3009'):
    sys.path.insert(0, str(ROOT / directory))
import facades_engine
import clad_engine
import frame_engine
from test_frame_solution_selection import example_selection
from probe_runtime import available, command, compile_probe

CHECKS = []

def need(value, reason):
    if not value:
        raise AssertionError(reason)
    CHECKS.append(reason)


def save(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + '\n', encoding='utf-8')


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def bounded_engine(engine, request):
    result = queue.Queue()
    def compute():
        try:
            result.put((True, engine.run(request)))
        except BaseException as error:
            result.put((False, error))
    threading.Thread(target=compute, daemon=True).start()
    try:
        ok, value = result.get(timeout=45)
    except queue.Empty:
        raise TimeoutError('Actual engine exceeded 45-second stage deadline')
    if not ok:
        raise value
    return value


def compile_native(build, extra_sources=(), extra_generated=(), extra_references=()):
    generated = []
    originals = []
    for name, module in (('CladdingQuantities', 'AClad'), ('FrameQuantities', 'AFrame')):
        source = ROOT / module / 'src' / (module + 'Plugin') / (name + '.cs')
        originals.append(source)
        raw = source.read_text(encoding='utf-8')
        marker = '        internal static QuantityReport BuildReport('
        need(raw.count(marker) == 1, 'producer extraction anchored uniquely: ' + name)
        prefix = '\n'.join(line for line in raw.splitlines() if line.startswith('using ') and 'Autodesk.' not in line)
        fields = ''
        if name == 'FrameQuantities':
            categories = re.findall(r'(?m)^\s*private static readonly string\[\] Categories = .*?;', raw)
            need(len(categories) == 1, 'frame category field extraction anchored uniquely')
            method_start = raw.index('        internal void SetProjectParameters(')
            method_end = raw.index('\n        // Expected visibility', method_start)
            fields = categories[0] + '\nprivate readonly QuantityReport report;\n' + \
                'internal FrameQuantities(QuantityReport existing) { report = existing; }\n' + raw[method_start:method_end]
        path = build / (name + 'Pure.cs')
        path.write_text(prefix + '\nnamespace ' + module + 'Plugin { internal sealed class ' + name + ' {\n' + fields + '\n' + raw[raw.index(marker):], encoding='utf-8')
        generated.append(path)
    source = ROOT / 'AFrame/src/AFramePlugin/FrameNodeCommand.cs'
    originals.append(source)
    raw = source.read_text(encoding='utf-8')
    start = raw.index('                if (current.ZoneId != saved.ZoneId')
    end = raw.index('\n                tr.Commit();', start)
    predicate = raw[start:end]
    need(predicate.count('E_NODE_SOURCE_STALE') == 1, 'actual command freshness predicate extraction anchored')
    path = build / 'NodeFreshnessPredicate.cs'
    path.write_text('''using System; using AFramePlugin; using Autodesk.AutoCAD.DatabaseServices;
internal static class NodeFreshnessPredicate {
internal sealed class Current { internal ObjectId ZoneId; internal string ZoneHandle,Fingerprint; internal FrameParameterResolution Resolution; }
internal sealed class Saved { internal ObjectId ZoneId; internal string ZoneHandle,ZoneFingerprint; internal FrameNodeSnapshot Snapshot; }
static void Fail(string code,string reason) { throw new InvalidOperationException(code+": "+reason); }
internal static bool Check(Current current,Saved saved) { try {
''' + predicate + '\nreturn true; } catch(InvalidOperationException) { return false; } } }\n', encoding='utf-8')
    generated.append(path)
    sources = [ROOT / 'Common' / (n + '.cs') for n in (
        'GeometryFingerprint', 'ZoneGeometryGuard', 'LayoutGeometryGuard', 'FacadeQuantities',
        'FacadeQuantityStore', 'FacadeQuantityStore.Frame', 'FacadeProjectParameterStore',
        'ManualQuantities', 'ManualQuantityGeometry', 'FacadeQuantityStore.Manual')]
    sources += [ROOT / 'AFrame/src/AFramePlugin' / (n + '.cs') for n in (
        'FrameSettings', 'FrameSolutionSelection', 'FrameProjectParameters', 'FrameNodeGeometry', 'FrameNodeDrawing', 'FrameMountingAssessment')]
    sources += [ROOT / 'Facades/src/AFacadesPlugin' / (n + '.cs') for n in (
        'CladdingTableData', 'FrameTableData', 'ConnectionTableData', 'QuantityTableIdentity',
        'QuantityTableView', 'QuantityTableView.Connections', 'QuantityXlsxFile', 'XlsxWriter')]
    sources += [ROOT / 'AClad/src/ACladPlugin/BondSettings.cs',
        ROOT / 'tools/quantities_3009/QuantityCadDoubles.cs', Path(__file__).with_name('LimitedPilotProbe.cs')]
    exe = build / 'LimitedPilotProbe.exe'
    sources += list(extra_sources); generated += list(extra_generated)
    compiled = compile_probe(sources + generated, ['System.Web.Extensions', 'System.IO.Compression', 'System.IO.Compression.FileSystem'] + list(extra_references), exe)
    (build / 'compile.log').write_text(compiled.stdout + compiled.stderr, encoding='utf-8')
    if compiled.returncode:
        print(compiled.stdout + compiled.stderr)
        raise RuntimeError('Native acceptance probe compilation failed')
    return exe, sources + originals, generated


def execute(exe, selection, out):
    """A reader thread gives every native event/exit a deadline; stderr drains separately."""
    events = queue.Queue()
    process = subprocess.Popen(command(exe) + [str(selection), str(out)], stdin=subprocess.PIPE,
        stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding='utf-8', bufsize=1)
    stderr = []
    def read_stdout():
        try:
            for line in process.stdout:
                events.put(line)
        finally:
            events.put(None)
    def read_stderr():
        for line in process.stderr:
            stderr.append(line)
    readers = [threading.Thread(target=read_stdout, daemon=True), threading.Thread(target=read_stderr, daemon=True)]
    for reader in readers:
        reader.start()
    observations = []
    try:
        while True:
            try:
                line = events.get(timeout=45)
            except queue.Empty:
                raise TimeoutError('Native stage produced no event within 45 seconds; process killed')
            if line is None:
                raise RuntimeError('Native process ended before completion: ' + ''.join(stderr))
            message = json.loads(line.lstrip('\ufeff'))
            if message['event_type'] == 'failure':
                raise RuntimeError(message['error'])
            if message['event_type'] == 'complete':
                need(process.wait(timeout=15) == 0, 'persistent native process exits cleanly')
                break
            need(message['event_type'] == 'engine', 'protocol only accepts explicit engine events')
            request = message['request']; before = copy.deepcopy(request)
            engine = {'zones': facades_engine, 'cladding': clad_engine, 'frame': frame_engine}[message['engine']]
            started = time.perf_counter(); response = bounded_engine(engine, request); elapsed = time.perf_counter() - started
            need(request == before, 'actual engine preserves exact captured request: ' + message['tag'])
            need(response.get('ok') is True, 'actual engine accepts: ' + message['tag'] + ': ' + str(response.get('error')))
            folder = out / 'engines' / message['tag']
            save(folder / 'request.json', request); save(folder / 'response.json', response)
            observations.append({'tag': message['tag'], 'engine': message['engine'], 'seconds': elapsed,
                'request_sha256': digest(folder / 'request.json'), 'response_sha256': digest(folder / 'response.json')})
            process.stdin.write(json.dumps(response, ensure_ascii=False, allow_nan=False) + '\n'); process.stdin.flush()
    finally:
        if process.poll() is None:
            process.kill()
        process.wait(timeout=15)
        for reader in readers:
            reader.join(timeout=5)
        (out / 'native_stderr.log').write_text(''.join(stderr), encoding='utf-8')
        save(out / 'engine_ledger.json', observations)
    return observations


def verify(out, observations):
    from openpyxl import load_workbook
    expected = json.loads((ROOT / 'docs/limited_pilot_0110/independent_oracle.json').read_text(encoding='utf-8'))['expected']
    frame_responses = []
    for item in observations:
        response = json.loads((out / 'engines' / item['tag'] / 'response.json').read_text(encoding='utf-8'))
        if item['engine'] == 'zones':
            need(len(response['zones']) == expected['zone_count'], 'independent zone count')
            need(abs(sum(z['report']['area_net_m2'] for z in response['zones']) - expected['zone_area_total_m2']) < 1e-10, 'independent wall area')
        elif item['engine'] == 'cladding':
            need(len(response['pieces']) == expected['tile_whole_each'], 'independent ATTILE piece count')
            need(sum(p['area'] for p in response['pieces']) == 2160000, 'independent ATTILE material area')
            need(response['per_zone'][0]['rows_y'] == expected['horizontal_joint_axes_mm'], 'independent horizontal joint axes')
        else:
            request = json.loads((out / 'engines' / item['tag'] / 'request.json').read_text(encoding='utf-8'))
            need(request['system']['bracket_step'] == request['system']['bracket_step_corner'] == 800 and
                 request['system']['name'] == 'Вектор-1' and request['parts'] == 'frame' and request['floors_y'] == [],
                 'normal actual manual800 pilot reaches engine without fallback or floor inputs')
            frame_responses.append(response)
            need(len(response['rails']) == 1 and len(response['brackets']) == 3, 'independent per-zone normal frame counts')
            rail = response['rails'][0]
            need(rail['x'] in expected['vertical_joint_axes_mm'] and rail['y0'] == 0 and rail['y1'] == 1820, 'independent guide position and length')
            need(sorted(b['y'] for b in response['brackets']) == [300, 1100, 1520], 'independent bracket positions')
            need(not any(response[k] for k in ('hrails', 'clamps', 'fittings')), 'no invented secondary components')
            need(response['solution_report']['physical_assignment'] == 'not_asserted', 'engineering product assignment remains unconfirmed')
    need(len(observations) == 9 and len(frame_responses) == 6, 'one zone run, two tile runs, six real frame computations including refused stale work')
    excel = []
    for tag in ('initial', 'project_reissue', 'zone_reissue'):
        for kind in ('cladding', 'frame', 'connections'):
            path = out / tag / (kind + '.xlsx')
            expected_rows = json.loads((out / tag / (kind + '_rows.json')).read_text(encoding='utf-8'))
            book = load_workbook(path, data_only=False); sheet = book.active
            for row_number, row in enumerate(expected_rows, 1):
                for column, value in enumerate(row, 1):
                    actual = sheet.cell(row_number, column).value
                    need(actual == value, 'actual XLSX matches model cell %s/%s/%s:%s' % (tag, kind, row_number, column))
            need(all(not isinstance(c.value, str) or len(c.value) <= 32767 for row in sheet for c in row), 'all exported XLSX cells fit Excel limits: ' + tag + '/' + kind)
            excel.append({'epoch': tag, 'view': kind, 'rows': sheet.max_row, 'columns': sheet.max_column, 'sha256': digest(path)})
            book.close()
    save(out / 'xlsx_review.json', {'status': 'PASS', 'files': excel, 'exact_model_values': True,
        'actual_file_publish_and_disposal_rollback': True, 'cad_transaction_rollback': 'NOT_RUN'})


def verify_trace(out):
    ledger = json.loads((out / 'ledger.json').read_text(encoding='utf-8'))
    epochs = {s['stage']: s for s in ledger if 'project' in s}
    initial = epochs['initial_fresh']; changed = epochs['project_changed_before_reissue']
    reissued = epochs['project_reissued']; overridden = epochs['zone_override_before_reissue']
    final = epochs['zone_reissued']
    zones = {z['zone_id']: z for z in initial['zones']}
    right = next(z['zone_id'] for z in zones.values() if z['xmin'] == 2410)
    left = next(z['zone_id'] for z in zones.values() if z['xmin'] == 0)
    need(left != right and len(zones) == 2, 'zone identities are matched through actual source contour coordinates')
    need(initial['project']['revision'] == 1 and changed['project']['revision'] == 2 and
         final['project']['revision'] == 2, 'one common project revision changes once')
    need(changed['project']['project_id'] == initial['project']['project_id'] == final['project']['project_id'], 'project identity survives both transitions')
    need(initial['project']['defaults']['geometry']['cladding_front_offset_mm'] == 230 and
         changed['project']['defaults']['geometry']['cladding_front_offset_mm'] == 240, 'actual stored declaration changes 230 to 240')
    initial_runs = {z['zone_id']: z['frame_run'] for z in initial['zones']}
    project_runs = {z['zone_id']: z['frame_run'] for z in reissued['zones']}
    final_runs = {z['zone_id']: z['frame_run'] for z in final['zones']}
    need({z['zone_id']: z['frame_run'] for z in changed['zones']} == initial_runs, 'global mutation retains old run identity until explicit reissue')
    need(all(project_runs[z] != initial_runs[z] for z in zones), 'both project reissues have new actual physical run IDs')
    need({z['zone_id']: z['frame_run'] for z in overridden['zones']} == project_runs, 'local mutation retains old run identity until explicit reissue')
    need(final_runs[left] == project_runs[left] and final_runs[right] != project_runs[right], 'local reissue changes only the right physical run')
    documents = set(); traces = []
    for tag, epoch in (('initial', initial), ('project_reissue', reissued), ('zone_reissue', final)):
        all_candidates = []
        for z in epoch['zones']:
            zid = z['zone_id']; effective_tag = 'project_reissue' if tag == 'zone_reissue' and zid == left else tag
            path = out / effective_tag / (zid + '_frame_report.json')
            report = json.loads(path.read_text(encoding='utf-8'))
            node = json.loads((out / effective_tag / (zid + '_node.json')).read_text(encoding='utf-8'))['snapshot']
            prepared = next(s for s in ledger if s['stage'] == 'prepare_frame' and s['tag'] == effective_tag and s['zone'] == zid)
            request = json.loads((out / 'engines' / (effective_tag + '_frame_' + zid) / 'request.json').read_text(encoding='utf-8'))
            response = json.loads((out / 'engines' / (effective_tag + '_frame_' + zid) / 'response.json').read_text(encoding='utf-8'))
            documents.add(report['document_id'])
            need(report['run_id'] == z['frame_run'] == z['source_run'], 'report, passport and persistent ledger preserve run identity')
            need(prepared['context'] == node['context'] == z['context'], 'captured context equals actual node snapshot and current ledger')
            snapshot = report['parameters']['project_parameters_snapshot']
            need(snapshot['project'] == prepared['context']['project'] == prepared['sources']['project_dependency'], 'stored report provenance equals dependency captured BEFORE engine')
            need(snapshot['zones'][0]['zone'] == prepared['context']['zone'], 'stored report preserves actual zone declaration revision')
            need(request['solution_selection'] == report['parameters']['solution_selection'] == node['selection'], 'node, engine request and physical inventory share exact resolved declaration')
            wanted_front = 230 if tag == 'initial' else (260 if tag == 'zone_reissue' and zid == right else 240)
            need(node['selection']['geometry']['cladding_front_offset_mm'] == wanted_front, 'project and local selection resolves to independently expected front')
            need(report['source_revisions']['owner_zone:' + z['owner']] == zid and report['zone_ids'] == [zid], 'actual stored quantity scope binds the same zone owner')
            elements = {e['element_id']: e for e in report['elements']}
            need(all(e['zone_ids'] == [zid] and e['product_id'] is None and len(e['cad_entities']) == 1 for e in elements.values()), 'physical IDs preserve scope and one managed CAD link without product assignment')
            need(len({e['cad_entities'][0]['handle'] for e in elements.values()}) == 4, 'one distinct managed physical handle per emitted item')
            passport = report['connection_passports'][0]
            need(passport['status'] == 'inventory_only' and passport['coverage']['strength'] == 'not_verified', 'accepted physical inventory does not grant engineering admission')
            candidates = []
            for member in passport['members']:
                need(elements[member['rail_element_id']]['role'] == 'rail', 'candidate member resolves to physical rail')
                for support in member['supports']:
                    for bracket in support['bracket_element_ids']:
                        need(elements[bracket]['role'] == 'bracket', 'candidate support resolves to physical bracket symbol')
                        candidates.append({'rail_element_id': member['rail_element_id'], 'bracket_element_id': bracket, 'offset_mm': support['offset_mm'], 'zone_id': zid})
            need(len(candidates) == 3, 'three geometric candidates per actual guide')
            need(set(elements) == {report['run_id'] + ':' + k + ':' + str(i) for k in ('rails', 'hrails', 'brackets', 'clamps', 'fittings') for i in range(len(response[k]))}, 'no physical array item is invented, lost or assigned a foreign run')
            all_candidates.extend(candidates)
            traces.append({'epoch': tag, 'zone_id': zid, 'owner': z['owner'], 'source_contour': z['source_contour'],
                'geometry_fingerprint': z['geometry'], 'run_id': report['run_id'], 'document_id': report['document_id'],
                'context': prepared['context'], 'source_snapshot': prepared['sources'], 'node_digest': node['snapshot_digest'],
                'physical': z['physical'], 'candidate_pairs': candidates, 'physical_uk': 'NOT_REPRESENTED', 'typed_interfaces': 'NOT_REPRESENTED'})
        need(len(all_candidates) == 6, 'independent total six geometric candidate pairs at every issued epoch')
    for path in (out / 'initial').glob('*_cladding_report.json'):
        report = json.loads(path.read_text(encoding='utf-8')); documents.add(report['document_id'])
    need(len(documents) == 1 and bool(next(iter(documents))), 'ATTILE and every frame generation belong to one persistent store document')
    stale = next(s for s in ledger if s['stage'] == 'stale_prepared_refused')
    need(stale['sources']['project_dependency']['revision'] == 1 and stale['counters']['dictionary_set'] == stale['counters']['xrecord_data_set'] == 0, 'rejected delayed response keeps old dependency and writes no records')
    save(out / 'trace.json', {'verification': 'PASS', 'engineering_assessment': 'BLOCKED', 'runs': traces,
        'identity_note': 'New producer generations have new physical IDs; unchanged geometry does not imply identical physical identities.',
        'node_scope': 'Pure snapshot/drawing plus source predicate only; actual command/store NOT_RUN'})


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--out', type=Path, required=True)
    args = ap.parse_args(); out = args.out.resolve(); out.mkdir(parents=True, exist_ok=True)
    if not available():
        print('BLOCKED: native .NET runtime/compiler unavailable'); return 2
    started = time.perf_counter(); build = out / 'build'; build.mkdir(exist_ok=True)
    exe, sources, generated = compile_native(build)
    compile_seconds = time.perf_counter() - started
    tracked = sources + [Path(__file__).resolve(), ROOT / 'tools/quantities_3009/probe_runtime.py',
        ROOT / 'docs/limited_pilot_0110/independent_oracle.json']
    for directory in ('AFrame/engine', 'AClad/engine', 'Facades/engine'):
        tracked.extend((ROOT / directory).glob('*.py'))
        tracked.extend((ROOT / directory).glob('*.json'))
    hashes = {str(p.relative_to(ROOT)): digest(p) for p in sorted(set(tracked))}
    selection = example_selection(); selection['geometry'].update(cladding_front_offset_mm=230, insulation_layers_mm=[100, 50])
    save(out / 'selection.json', selection)
    tick = time.perf_counter(); observations = execute(exe, out / 'selection.json', out); native_seconds = time.perf_counter() - tick
    verify(out, observations)
    verify_trace(out)
    native = json.loads((out / 'native_result.json').read_text(encoding='utf-8'))
    need(all(digest(ROOT / name) == value for name, value in hashes.items()), 'all bound sources remain unchanged during run')
    artifacts = {str(p.relative_to(out)): digest(p) for p in sorted(out.rglob('*.json'))
        if build not in p.parents and p.name != 'manifest.json'}
    manifest = {'schema': 'limited_pilot_acceptance/1', 'status': 'PASS', 'verification': 'PASS',
        'engineering_assessment': 'BLOCKED', 'full_P5': 'NOT_RUN', 'full_2V': 'BLOCKED', 'calculation_3A': 'BLOCKED',
        'live_autocad_checked': False, 'native_checks': native['checks'], 'python_checks': len(CHECKS),
        'checks': native['checks'] + len(CHECKS), 'engine_invocations': len(observations), 'xlsx_files': 9,
        'source_sha256': hashes, 'generated_extraction_sha256': {str(p.relative_to(out)): digest(p) for p in generated},
        'artifact_sha256': artifacts,
        'sources_unchanged_during_run': True, 'assertions': CHECKS,
        'timings': {'compile_seconds': compile_seconds, 'persistent_pipeline_seconds': native_seconds, 'scope': 'managed acceptance harness only; no live AutoCAD speed claim'},
        'limits': ['CAD entities and transactions are QuantityCadDoubles; Commit/Dispose are no-ops',
            'No DWG file fabricated; native save/open/Undo/COPY/render NOT_RUN',
            'Node pure geometry/drawing plus verbatim source predicate; actual ATFNODE command and node store NOT_RUN',
            'Geometry materialization/CAD-link bridge substitutes command drawing adapters; no whole UI execution',
            'Happy-flow bridge materializes and erases before StoreFrame; this is not a product prewrite/rollback proof',
            'Actual XLSX file publication/disposal rollback exercised; CAD rollback not tested']}
    save(out / 'manifest.json', manifest)
    print('Limited pilot acceptance:', manifest['checks'], 'PASS; mounting/3A BLOCKED; full P5 NOT_RUN')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
