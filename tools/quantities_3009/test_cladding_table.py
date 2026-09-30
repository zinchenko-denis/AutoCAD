"""Execute real table model and XLSX writer, inspect exported numbers with openpyxl.

Uses synthetic public rows only; no AutoCAD execution is implied.
"""
import argparse
import hashlib
import json
from pathlib import Path
import subprocess
import tempfile
from probe_runtime import available, command, compile_probe

ROOT = Path(__file__).resolve().parents[2]


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--out', type=Path)
    args = ap.parse_args()
    if not available():
        print('BLOCKED: native .NET test compiler/runtime unavailable; no C# execution claimed')
        return 2
    out = (args.out or Path(tempfile.mkdtemp(prefix='cladding_table_'))).resolve()
    out.mkdir(parents=True, exist_ok=True)
    sources = [ROOT / 'tools/quantities_3009/CladdingTableProbe.cs',
               ROOT / 'Common/FacadeQuantities.cs',
               ROOT / 'Facades/src/AFacadesPlugin/CladdingTableData.cs',
               ROOT / 'Facades/src/AFacadesPlugin/QuantityXlsxFile.cs',
               ROOT / 'Facades/src/AFacadesPlugin/XlsxWriter.cs']
    executable = out / 'CladdingTableProbe.exe'
    compiled = compile_probe(sources, ['System.IO.Compression', 'System.IO.Compression.FileSystem', 'System.Web.Extensions', 'System.Xml'], executable)
    (out / 'compile.log').write_text(compiled.stdout + compiled.stderr, encoding='utf-8')
    if compiled.returncode:
        print(compiled.stdout + compiled.stderr)
        return 1
    subprocess.run(command(executable) + [str(out)], check=True)

    from openpyxl import load_workbook
    exact = load_workbook(out / 'numeric_exact.xlsx')['Числа']
    legacy = load_workbook(out / 'numeric_legacy.xlsx')['Числа']
    checks = 0

    def check(value, message):
        nonlocal checks
        checks += 1
        assert value, message

    check(exact['A1'].value == 0.1234567891234567, 'small area loses precision')
    check(exact['B1'].value == 123456789.12345678, 'large value loses precision')
    check(exact['A2'].value == 0.0000000123456789, 'tiny positive value becomes zero')
    check(exact['C1'].value is None, 'missing value becomes zero')
    check(exact['D1'].data_type == 's' and exact['D1'].value == '=SUM(1,2)', 'text interpreted as formula')
    check(exact['E1'].value == '  A & <B>  ', 'XML or whitespace corruption')
    check(exact['C2'].value == 'не определена', 'unknown value is not explicit')
    check(legacy['A1'].value == 0.123 and legacy['A2'].value == 0, 'existing writer numeric contract changed')

    sheet = load_workbook(out / 'cladding_exact.xlsx')['Облицовка']
    rows = list(sheet.values)
    installed = next(row for row in rows if row[0] == 'А' and row[1] == 'Керамогранит & <серия>')
    subtotal = next(row for row in rows if row[0] == 'ИТОГО')
    check(installed[7] == 2 and installed[8] == 0.1234567891234567, 'installed values differ from renderer')
    check(subtotal[7] == 5 and subtotal[8] == 'не определена', 'subtotal is rounded or includes cutting')
    check(any(row[5] == 'заготовки' and row[7] == 17 for row in rows if len(row) > 7), 'cutting section missing')
    check(sum(str(row[0]).startswith('Неопределённость:') for row in rows) == 2, 'warnings lost or duplicated')
    check(rows[-1][0] == 'Примечание пользователя:   Примечание & <узел> =1+1  ', 'user note lost')
    check(all(cell.data_type != 'f' for row in sheet for cell in row), 'unexpected formula in quantity workbook')
    manifest = {'status': 'PASS', 'xlsx_checks': checks, 'live_autocad_checked': False,
                'source_sha256': {str(path.relative_to(ROOT)): hashlib.sha256(path.read_bytes()).hexdigest()
                                  for path in sources}}
    (out / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'CladdingTable XLSX checks: {checks}; PASS')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
